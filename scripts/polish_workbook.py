import copy, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

SRC, DST = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)

NAVY, BLUE, GREEN, GREY = "E4003A", "C8102E", "F04E23", "6B6B6B"  # crimson palette
BAND = "FFF5F7"
SECTION = "FCE4EA"
thin = Side(style="thin", color="EBD3D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor=NAVY)
FONT = "Arial"


def arial(cell, size=None, bold=None, color=None, italic=None, underline=None):
    f = cell.font
    col = color if color is not None else (f.color.rgb if f.color is not None and isinstance(f.color.rgb, str) else None)
    cell.font = Font(name=FONT,
                     size=size if size else (f.sz if f.sz and f.sz != 11 else 10),
                     bold=f.b if bold is None else bold,
                     italic=f.i if italic is None else italic,
                     underline=f.u if underline is None else underline,
                     color=col)


def has_fill(c):
    return c.fill is not None and c.fill.fill_type == "solid" and c.fill.fgColor.rgb not in ("00000000", None)


def base_font(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None or has_fill(c):
                arial(c)


# ---------- generic table styling ----------
WIDTH = {
    "route": 8, "row_type": 11, "parent": 24, "brand": 28, "category": 22, "year_start": 9,
    "consideration_inr_cr": 15, "stake_pct": 9, "deal_structure": 18, "revenue_inr_cr": 14,
    "revenue_fy": 9, "revenue_basis": 30, "revenue_multiple": 11, "what_was_acquired": 18,
    "vocab_rationale": 46, "status_today": 32, "status_evidence": 52, "base_rate_flag": 20,
    "source_url": 26, "source_ids": 16, "date_accessed": 12, "confidence": 12, "notes": 70,
    "field": 18, "value": 24, "tie_break_rank": 10, "inclusion_rule": 70,
    "source_id": 10, "url": 34, "title": 46, "publisher": 26, "tier": 6, "date_published": 13,
    "claim_id": 9, "item": 28, "value_recalled": 40, "value_verified": 46, "match": 9, "how_caught": 52,
    "fy": 7, "mentions": 10, "pages_with_mentions": 11, "sample_page": 9, "sample_quote": 70,
    "ar_url": 26, "status": 14, "definition": 46, "numerator": 11, "denominator": 12, "rate": 9,
    "wilson_95": 11, "rate_if_stable_counts": 13, "rows_included": 60,
    "table": 7, "v1_demand": 9, "v2_distance": 10, "v3_window": 9, "v5_founder": 10, "score": 8,
    "side": 12, "hit": 9,
    "assumption_id": 9, "parameter": 26, "value_low": 11, "value_base": 11, "value_high": 11,
    "unit": 34, "build_favourable": 11, "Bull": 12, "Base": 12, "Bear": 12, "rationale": 60,
}
WRAP = {"revenue_basis", "vocab_rationale", "status_today", "status_evidence", "notes", "inclusion_rule",
        "title", "value_recalled", "value_verified", "how_caught", "sample_quote", "definition",
        "rows_included", "unit", "rationale", "parent", "brand", "item", "source_ids", "category"}
CENTER = {"route", "row_type", "year_start", "stake_pct", "revenue_fy", "confidence", "tier", "fy",
          "mentions", "pages_with_mentions", "sample_page", "match", "claim_id", "source_id", "table",
          "v1_demand", "v2_distance", "v3_window", "v5_founder", "score", "hit", "side", "tie_break_rank",
          "numerator", "denominator", "rate", "wilson_95", "rate_if_stable_counts", "date_accessed",
          "date_published", "assumption_id", "build_favourable"}


def style_table(ws, last_row=None, row_height=None, tab_color=BLUE):
    ws.sheet_properties.tabColor = tab_color
    last_row = last_row or ws.max_row
    ncol = ws.max_column
    hdr = [ws.cell(1, i).value for i in range(1, ncol + 1)]
    ws.row_dimensions[1].height = 34
    for i, h in enumerate(hdr, 1):
        c = ws.cell(1, i)
        if h is None:
            continue
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(h, 14)
    for r in range(2, last_row + 1):
        if row_height:
            ws.row_dimensions[r].height = row_height
        for i, h in enumerate(hdr, 1):
            c = ws.cell(r, i)
            if h is None:
                continue
            arial(c)
            c.border = BORDER
            c.alignment = Alignment(vertical="top",
                                    horizontal="center" if h in CENTER else None,
                                    wrap_text=h in WRAP)
            if r % 2 == 0 and not has_fill(c):
                c.fill = PatternFill("solid", fgColor=BAND)
            if h in ("source_url", "url", "ar_url") and c.value:
                c.font = Font(name=FONT, size=10, color="0563C1", underline="single")
            if h == "stake_pct" and isinstance(c.value, (int, float)):
                c.number_format = '0.0"%"'
            if h in ("year_start", "numerator", "denominator", "mentions"):
                c.number_format = "0"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"


# ---------- data and log tabs ----------
style_table(wb["Tab_A_Buy"], row_height=96)
style_table(wb["Tab_B_Build"], row_height=96)
style_table(wb["Vocabulary"], tab_color=BLUE)
style_table(wb["Source_Log"], tab_color=GREY)
style_table(wb["Claims_Log"], row_height=48, tab_color=GREY)
style_table(wb["AR_Mentions"], row_height=48, tab_color=GREY)
style_table(wb["Framework_Scores"], tab_color=GREEN)

# Base_Rates: make rate a live formula
br = wb["Base_Rates"]
for r in range(2, br.max_row + 1):
    if isinstance(br.cell(r, 2).value, (int, float)) and isinstance(br.cell(r, 3).value, (int, float)):
        br.cell(r, 4).value = f"=IFERROR(B{r}/C{r},\"\")"
        br.cell(r, 4).number_format = "0%"
style_table(br, row_height=120, tab_color=GREEN)
for r in range(2, br.max_row + 1):
    br.cell(r, 4).font = Font(name=FONT, size=10, bold=True)

# Assumptions: table rows 1-26, notes below
a = wb["Assumptions"]
style_table(a, last_row=26, row_height=60, tab_color=GREEN)
for r in range(27, a.max_row + 1):
    for c in a[r]:
        if c.value is not None:
            arial(c)
            c.alignment = Alignment(vertical="top", wrap_text=False)
a.cell(28, 1).font = Font(name=FONT, size=10, bold=True, color=NAVY)
for r in range(29, 34):
    a.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=False)

# Conflicts
cf = wb["Conflicts"]
cf.sheet_properties.tabColor = GREY
base_font(cf)
for row in cf.iter_rows():
    for c in row:
        if c.value is not None:
            c.alignment = Alignment(vertical="top", wrap_text=c.column > 1 or not has_fill(c) and len(str(c.value)) < 60)
            if c.value in ("Figure", "Source (tier)", "What it measures", "Verdict"):
                c.font = Font(name=FONT, size=10, bold=True, color=NAVY)
    if row[0].value and has_fill(row[0]):
        for k in range(1, 6):
            cf.cell(row[0].row, k).fill = PatternFill("solid", fgColor=SECTION)
        row[0].font = Font(name=FONT, size=11, bold=True, color=NAVY)
        row[0].alignment = Alignment(vertical="center", wrap_text=False)
cf["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
cf.page_setup.orientation = "landscape"
cf.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
cf.page_setup.fitToWidth = 1
cf.page_setup.fitToHeight = 0
for k, w in zip("ABCDE", (34, 30, 48, 60, 40)):
    cf.column_dimensions[k].width = w

# ---------- model tabs ----------
def style_model(ws, label_col):
    ws.sheet_properties.tabColor = GREEN
    base_font(ws)
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="595959")
    for row in ws.iter_rows():
        a_val = row[0].value
        lab = row[label_col - 1].value if len(row) >= label_col else None
        if isinstance(a_val, str) and (a_val.startswith("Scenario") or a_val.startswith("Grid block")
                                        or a_val.startswith("Key outputs") or a_val.startswith("Calculation")
                                        or a_val.startswith("Grid:") or a_val.startswith("Tornado")
                                        or a_val.startswith("Monte")):
            for k in range(1, 14):
                c = ws.cell(row[0].row, k)
                c.fill = PatternFill("solid", fgColor=NAVY if a_val.startswith(("Scenario", "Grid block")) else SECTION)
            row[0].font = Font(name=FONT, size=11, bold=True,
                               color="FFFFFF" if a_val.startswith(("Scenario", "Grid block")) else NAVY)
        if isinstance(lab, str) and lab.startswith("Year (Y0"):
            for k in range(label_col, 14):
                c = ws.cell(row[0].row, k)
                c.font = Font(name=FONT, size=10, bold=True)
                c.fill = PatternFill("solid", fgColor=SECTION)
                c.border = Border(bottom=Side(style="thin", color=NAVY))
        if isinstance(lab, str) and lab.upper().startswith(("BREAKEVEN", "BUILD-EQUIVALENT", "RISK-ADJUSTED", "COST OF DELAY")):
            ws.cell(row[0].row, label_col).font = Font(name=FONT, size=10, bold=True)
            ws.cell(row[0].row, label_col + 1).font = Font(name=FONT, size=10, bold=True)
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


bm = wb["Build_Model"]
style_model(bm, 2)
bm.column_dimensions["B"].width = 66

s = wb["Sensitivity"]
style_model(s, 2)
pretty = {
    "years_from_decision": "Years from decision to ₹347 Cr",
    "pv_burn_to_scale": "PV of cash burn to scale",
    "p_success": "P(success)",
    "risk_adjusted_build_cost": "Risk-adjusted build cost (PV burn ÷ p)",
    "cost_of_delay": "Cost of delay",
    "build_equivalent_cost": "Build-equivalent cost",
    "buy_all_in_100pct": "Buy all-in (100%)",
    "breakeven_headline_price": "Breakeven price for 90.5%",
    "headroom_vs_price": "Headroom vs price paid",
}
for r in range(5, 14):
    v = s.cell(r, 1).value
    if v in pretty:
        s.cell(r, 1).value = pretty[v]
s.column_dimensions["A"].width = 40
s.column_dimensions["B"].width = 44
for k in range(2, 5):
    s.cell(4, k).font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    s.cell(4, k).fill = HDR_FILL
    s.cell(4, k).alignment = Alignment(horizontal="center")
for r in (12,):
    for k in range(1, 5):
        s.cell(r, k).font = Font(name=FONT, size=10, bold=True)
for k in range(1, 8):
    c = s.cell(17, k)
    c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    c.fill = HDR_FILL
    c.alignment = Alignment(horizontal="center")
for r in range(18, 24):
    s.cell(r, 1).font = Font(name=FONT, size=10, bold=True)
    s.cell(r, 1).alignment = Alignment(horizontal="center")
    for k in range(1, 8):
        s.cell(r, k).border = BORDER
s.cell(16, 1).font = Font(name=FONT, size=10, italic=True, color="595959")

# ---------- README cover ----------
old = wb["README"]
idx = wb.sheetnames.index("README")
lines = [old.cell(r, 1).value for r in range(1, old.max_row + 1)]
wb.remove(old)
rd = wb.create_sheet("README", idx)
wb.active = idx
rd.sheet_properties.tabColor = NAVY
rd.sheet_view.showGridLines = False
for k, w in zip("ABCDE", (2, 46, 16, 80, 2)):
    rd.column_dimensions[k].width = w

def put(r, col, val, **kw):
    c = rd.cell(r, col, val)
    c.font = Font(name=FONT, size=kw.get("size", 10), bold=kw.get("bold", False), italic=kw.get("italic", False),
                  color=kw.get("color", "000000"), underline=kw.get("underline"))
    c.alignment = Alignment(vertical="center", wrap_text=kw.get("wrap", False), horizontal=kw.get("h"))
    if kw.get("fill"):
        c.fill = PatternFill("solid", fgColor=kw["fill"])
    if kw.get("fmt"):
        c.number_format = kw["fmt"]
    return c

def band(r, text):
    for k in range(2, 5):
        rd.cell(r, k).fill = HDR_FILL
    put(r, 2, text, size=11, bold=True, color="FFFFFF", fill=NAVY)
    rd.row_dimensions[r].height = 22

for k in range(1, 6):
    rd.cell(1, k).fill = PatternFill("solid", fgColor=NAVY)
rd.row_dimensions[1].height = 8
put(2, 2, "Buy or Build", size=20, bold=True, color=NAVY)
put(3, 2, "Consolidation in Indian consumer brands: evidence workbook", size=12, color="404040")
put(4, 2, "Consumer Brands Research · Prepared by Priyanshi Paul · October 2026",
    size=9, italic=True, color="7F7F7F")
rd.row_dimensions[2].height = 30

band(6, "Headline results (live, linked to the model)")
CR = '"₹"#,##0" Cr"'
kpis = [
    ("Price paid at completion, 90.5% of Minimalist", "=Sensitivity!B14", CR, "HUL completion filing, 2025-04-21 (see Conflicts #1)"),
    ("Risk-adjusted build-equivalent cost (base)", "=Sensitivity!B10", CR, "PV of burn ÷ P(success) + cost of delay"),
    ("Buy all-in cost, 100% (base)", "=Sensitivity!B11", CR, "Price + last 9.5% + transaction + integration − cash acquired"),
    ("Breakeven price for 90.5% (base)", "=Sensitivity!B12", CR, "Above this, buying stops making sense"),
    ("Headroom vs price paid (base)", "=Sensitivity!B13", CR, "Positive = buying was cheaper"),
    ("P(success) of an in-house build (base)", "=Sensitivity!B7", "0%", "Laplace estimate from 0 of 7 FMCG digital-first builds (Base_Rates c)"),
    ("Years from decision to ₹347 Cr (base)", "=Sensitivity!B5", '0" years"', "Bull 5 · Bear 9"),
]
r = 7
for lab, f, fmt, note in kpis:
    put(r, 2, lab)
    put(r, 3, f, bold=True, color=NAVY, fmt=fmt, h="right")
    put(r, 4, note, size=9, italic=True, color="595959")
    for k in range(2, 5):
        rd.cell(r, k).border = Border(bottom=Side(style="thin", color="EBD3D9"))
    rd.row_dimensions[r].height = 18
    r += 1

r += 1
band(r, "Dataset at a glance"); r += 1
counts = [
    ("Acquisitions (Tab A)", "=COUNTA(Tab_A_Buy!A2:A1000)", "Listed parents, 2020 onward"),
    ("In-house builds (Tab B)", "=COUNTA(Tab_B_Build!A2:A1000)", "Core (2020+) and calibration (pre-2020) rows"),
    ("Sources logged", "=COUNTA(Source_Log!A2:A1000)", "Exchange filings and annual reports first"),
    ("AI claims checked", "=COUNTA(Claims_Log!A2:A1000)", "Every figure recalled from model memory, verified against a source"),
    ("Annual-report brand mentions tracked", "=COUNTA(AR_Mentions!A2:A1000)", "Brand × parent × FY rows, FY20–FY26"),
]
for lab, f, note in counts:
    put(r, 2, lab)
    put(r, 3, f, bold=True, color=NAVY, fmt="#,##0", h="right")
    put(r, 4, note, size=9, italic=True, color="595959")
    for k in range(2, 5):
        rd.cell(r, k).border = Border(bottom=Side(style="thin", color="EBD3D9"))
    r += 1

r += 1
band(r, "Tabs (click to open)"); r += 1
tab_lines = [l for l in lines if isinstance(l, str) and ":" in l and l.split(":")[0] in wb.sheetnames]
for l in tab_lines:
    name, desc = l.split(":", 1)
    c = put(r, 2, name, color="0563C1", underline="single")
    c.hyperlink = f"#'{name}'!A1"
    put(r, 3, desc.strip(), wrap=True); rd.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    rd.row_dimensions[r].height = 30
    for k in range(2, 5):
        rd.cell(r, k).border = Border(bottom=Side(style="thin", color="EBD3D9"))
    r += 1

r += 1
band(r, "Source tiers"); r += 1
for l in lines:
    if isinstance(l, str) and l[:2] in ("T1", "T2", "T3", "T4") and l[2] == ":":
        put(r, 2, l[:2], bold=True)
        put(r, 3, l[3:].strip(), wrap=True); rd.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        r += 1

r += 1
band(r, "Confidence rules"); r += 1
conf = [("High", "C6EFCE", "A T1 source with a verbatim quote."),
        ("Medium", "FFEB9C", "A T2 source, or two independent T3 sources that agree."),
        ("Low", "FFC7CE", "A single T3 source, or a derived or inferred value."),
        ("UNCONFIRMED", "FFC7CE", "Value left blank with the reason in notes; never estimated silently.")]
for lab, col, txt in conf:
    put(r, 2, lab, bold=True, fill=col)
    put(r, 3, txt, wrap=True); rd.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    r += 1

r += 1
band(r, "Conventions"); r += 1
put(r, 2, "1,000", color="0000FF"); put(r, 3, "Blue font = hard-coded input (edit these on Assumptions)"); r += 1
put(r, 2, "formula output", fill="FFF2CC"); put(r, 3, "Yellow fill = live formula output"); r += 1
put(r, 2, "₹ Cr · ISO dates"); put(r, 3, "All money in ₹ crore; FY24 = Apr 2023 to Mar 2024; revenue basis recorded per row"); r += 1
put(r, 2, "Conflicts"); put(r, 3, "Resolved in favour of the most authoritative source; every figure and what it measures is kept"); r += 1

rd.page_setup.orientation = "portrait"
rd.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
rd.page_setup.fitToWidth = 1
rd.page_setup.fitToHeight = 0

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 100
wb.active = 0
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = ws.title == "README"
wb.save(DST)
print("saved", DST)
