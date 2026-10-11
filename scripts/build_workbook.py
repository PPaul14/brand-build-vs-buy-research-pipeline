"""Build output/BuyOrBuild_Workbook.xlsx from data/*.csv (openpyxl; Google Sheets compatible).

Plain formulas only: no data tables, array formulas, macros or external links. revenue_multiple, the Build_Model
(Bull/Base/Bear) and the Sensitivity grid are live formulas that read the labelled cells on the Assumptions tab.
The tornado and Monte Carlo blocks are static outputs of scripts/build_model.py (10,000 draws cannot be a formula).

  python scripts/build_workbook.py

After building, run scripts/polish_workbook.py for styling.
"""
import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "output" / "BuyOrBuild_Workbook.xlsx"

MONEY = '"₹"#,##0.00" Cr"'
MONEY0 = '"₹"#,##0" Cr"'
MULT = '0.00"x"'
PCT = "0.0%"
BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HDR_FONT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
OUT_FILL = PatternFill("solid", fgColor="FFF2CC")
GREEN = PatternFill("solid", fgColor="C6EFCE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
LINK = Font(color="0563C1", underline="single")
THIN = Border(top=Side(style="thin", color="808080"))
WRAP = Alignment(wrap_text=True, vertical="top")

# Parameters stored as percent numbers in assumptions.csv; written to the sheet as fractions.
PCT_PARAMS = {"gross_margin_early", "gross_margin_scale", "ap_pct_launch", "ap_pct_growth", "ap_pct_scale",
              "var_opex_pct", "org_cost_pct", "rnd_pct", "discount_rate", "category_growth", "p_success",
              "steady_ebitda_margin", "terminal_growth"}
# Same mapping as BUILD_FAVOURABLE in build_model.py.
FAV = {
    "years_to_scale": "low", "y1_revenue": "high", "gross_margin_early": "high", "gross_margin_scale": "high",
    "ap_pct_launch": "low", "ap_pct_growth": "low", "ap_pct_scale": "low", "var_opex_pct": "low",
    "org_cost_floor": "low", "org_cost_pct": "low", "rnd_floor": "low", "rnd_pct": "low", "wc_days": "low",
    "category_growth": "low", "p_success": "high", "steady_ebitda_margin": "low", "acquired_rev_y0": "low",
    "integration_cost": "high", "remaining_stake_cost": "high", "cash_acquired": "low",
}
MONEY_COLS = {"consideration_inr_cr", "revenue_inr_cr"}


def read(name):
    with (DATA / name).open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


NUM_RE = re.compile(r"-?\d+(\.\d+)?")


def num(x):
    if x is None:
        return x
    x = ILLEGAL_CHARACTERS_RE.sub("", x)
    if not NUM_RE.fullmatch(x.strip()):
        return x
    return float(x) if "." in x else int(x)


def header(ws, cols, row=1):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row, j, c)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, w):
    for j, x in enumerate(w, 1):
        ws.column_dimensions[L(j)].width = x


def finish(ws, ncols, nrows, freeze="A2", hdr_row=1):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A{hdr_row}:{L(ncols)}{max(nrows, hdr_row + 1)}"


def confidence_rules(ws, col, first, last):
    rng = f"{col}{first}:{col}{last}"
    c = f"{col}{first}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{c}="High"'], fill=GREEN))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{c}="Medium"'], fill=AMBER))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'OR({c}="Low",ISNUMBER(SEARCH("UNCONFIRMED",{c})))'], fill=RED))


def link(cell, url):
    url = (url or "").split(";")[0].strip()
    if url.startswith("http"):
        cell.hyperlink = url
        cell.font = LINK


def dump(ws, rows, money=(), wide=(), url_cols=(), conf_col=None):
    cols = list(rows[0].keys())
    header(ws, cols)
    for i, r in enumerate(rows, 2):
        for j, c in enumerate(cols, 1):
            cell = ws.cell(i, j, num(r[c]))
            if c in money:
                cell.number_format = MONEY
            if c in url_cols:
                link(cell, r[c])
    widths(ws, [45 if c in wide else 16 for c in cols])
    finish(ws, len(cols), len(rows) + 1)
    if conf_col and conf_col in cols:
        confidence_rules(ws, L(cols.index(conf_col) + 1), 2, len(rows) + 1)
    return cols


# ---------------------------------------------------------------- deal / build tabs
def deal_tab(ws, rows):
    cols = list(rows[0].keys())
    header(ws, cols)
    ix = {c: L(j) for j, c in enumerate(cols, 1)}
    for i, r in enumerate(rows, 2):
        for j, c in enumerate(cols, 1):
            cell = ws.cell(i, j)
            if c == "revenue_multiple":
                k, s, v = (f"{ix[x]}{i}" for x in ("consideration_inr_cr", "stake_pct", "revenue_inr_cr"))
                cell.value = f'=IF(AND(ISNUMBER({k}),ISNUMBER({s}),ISNUMBER({v}),N({v})>0),{k}/({s}/100)/{v},"")'
                cell.number_format = MULT
                cell.fill = OUT_FILL
                continue
            cell.value = r[c] if c == "source_url" else num(r[c])
            if c in MONEY_COLS:
                cell.number_format = MONEY
            if c in ("consideration_inr_cr", "stake_pct", "revenue_inr_cr"):
                cell.font = BLUE
            if c == "source_url":
                link(cell, r[c])
            if c in ("vocab_rationale", "status_today", "status_evidence", "notes"):
                cell.alignment = WRAP
    wide = {"brand": 32, "vocab_rationale": 45, "status_today": 40, "status_evidence": 45, "notes": 60,
            "source_url": 30, "parent": 22}
    widths(ws, [wide.get(c, 13) for c in cols])
    finish(ws, len(cols), len(rows) + 1, freeze="E2")
    confidence_rules(ws, ix["confidence"], 2, len(rows) + 1)


# ---------------------------------------------------------------- assumptions
ASM_COLS = ["assumption_id", "parameter", "value_low", "value_base", "value_high", "unit", "build_favourable",
            "Bull", "Base", "Bear", "source_ids", "rationale", "confidence"]
CONSTANTS = [  # (key, label, value, fmt, source note) -- constants hard-coded in build_model.py
    ("headline_price", "Headline price paid for 90.5% (cash at completion)", 2706.45, MONEY, "S002 HUL completion letter"),
    ("m0", "Minimalist EBITDA margin at Y0 (FY25: ₹18 Cr / ₹514.8 Cr)", 0.035, PCT, "S209"),
    ("psg1", "Built brand growth, 1st year after scale (Minimalist FY25)", 0.48, PCT, "S209"),
    ("psg2", "Built brand growth, 2nd year after scale (Minimalist FY26 annualised)", 0.34, PCT, "S134"),
    ("ramp", "Years for EBITDA margin to reach steady state", 5, "0", "Modelling choice"),
]


def assumptions_tab(ws, rows):
    """Returns {scenario: {param: absolute ref}} plus constants under key 'const'."""
    header(ws, ASM_COLS)
    refs = {"Bull": {}, "Base": {}, "Bear": {}, "const": {}}
    for i, r in enumerate(rows, 2):
        p = r["parameter"]
        scale = 0.01 if p in PCT_PARAMS else 1
        fmt = PCT if p in PCT_PARAMS else ("0" if p in ("years_to_scale", "failed_attempt_years", "wc_days")
                                            else MONEY if "Rs Cr" in r["unit"] else "0.00")
        vals = dict(r, build_favourable=FAV.get(p, "base"))
        for j, c in enumerate(ASM_COLS, 1):
            cell = ws.cell(i, j)
            if c in ("value_low", "value_base", "value_high"):
                cell.value = float(r[c]) * scale
                cell.font, cell.number_format = BLUE, fmt
            elif c in ("Bull", "Base", "Bear"):
                lo, ba, hi, fv = f"$C{i}", f"$D{i}", f"$E{i}", f"$G{i}"
                cell.value = {"Base": f"={ba}",
                              "Bull": f'=IF({fv}="low",{lo},IF({fv}="high",{hi},{ba}))',
                              "Bear": f'=IF({fv}="low",{hi},IF({fv}="high",{lo},{ba}))'}[c]
                cell.number_format, cell.fill = fmt, OUT_FILL
                refs[c][p] = f"Assumptions!${L(j)}${i}"
            elif c == "build_favourable":
                cell.value = vals[c]
                cell.font = BLUE
            else:
                cell.value = num(vals[c]) if c != "source_ids" else vals[c]
                if c == "rationale":
                    cell.alignment = WRAP
    n = len(rows) + 1
    confidence_rules(ws, "M", 2, n)
    r0 = n + 2
    ws.cell(r0, 1, "Model constants (hard-coded in scripts/build_model.py)").font = BOLD
    ws.cell(r0, 1).fill = SEC_FILL
    for k, (key, lab, val, fmt, src) in enumerate(CONSTANTS, r0 + 1):
        ws.cell(k, 1, key)
        ws.cell(k, 2, lab)
        c = ws.cell(k, 4, val)
        c.font, c.number_format = BLUE, fmt
        ws.cell(k, 11, src)
        refs["const"][key] = f"Assumptions!$D${k}"
    note = r0 + len(CONSTANTS) + 2
    ws.cell(note, 1, "Blue = input. Yellow = formula. % drivers are stored as fractions (60% = 0.60). "
                     "Bull = every driver at its build-favourable end, Bear = the opposite; valuation inputs "
                     "(discount rate, terminal growth, target, transaction costs, failed-attempt years) stay at base.")
    widths(ws, [8, 24, 11, 11, 11, 40, 12, 11, 11, 11, 20, 70, 11])
    finish(ws, len(ASM_COLS), n)
    return refs


# ---------------------------------------------------------------- build model block (one scenario)
YEARS = list(range(11))
FIRST_YC = 3  # column C = year 0


def model_block(ws, top, title, P, C):
    """Write a Y0-Y10 build P&L + valuation block starting at row `top`. P = param refs, C = constant refs.
    Returns (summary refs dict, next free row). Mirrors build_pnl/bought_path/evaluate in build_model.py."""
    ycol = [L(FIRST_YC + t) for t in YEARS]
    yfirst, ylast = ycol[0], ycol[-1]
    ws.cell(top, 1, title).font = Font(bold=True, size=12)
    ws.cell(top, 1).fill = SEC_FILL
    r = top + 1
    # scalars
    sc = {}
    for key, lab, f, fmt in [
        ("T", "Years to scale T (rounded)", f"=ROUND({P['years_to_scale']},0)", "0"),
        ("g", "Growth to scale g = (target / Y1 revenue)^(1/(T-1))",
         None, PCT),
        ("k", "Failed-attempt years (int)", f"=INT({P['failed_attempt_years']})", "0"),
    ]:
        ws.cell(r, 2, lab)
        sc[key] = f"$C${r}"
        if key == "g":
            f = f"=({P['target_revenue']}/{P['y1_revenue']})^(1/({sc['T']}-1))"
        c = ws.cell(r, 3, f)
        c.number_format = fmt
        r += 1
    r += 1
    ws.cell(r, 2, "Year (Y0 = FY26)").font = BOLD
    for t in YEARS:
        ws.cell(r, FIRST_YC + t, t).font = BOLD
    row_year = r
    r += 1
    lines = ["revenue", "gross_margin", "ap_rate", "gross_profit", "marketing", "rnd", "org_cost", "other_opex_pre",
             "ebitda_pre", "other_opex", "ebitda", "wc", "wc_investment", "fcf", "cash_burn", "cumulative_cash",
             "discount_factor", "bought_revenue", "bought_ebitda"]
    R = {name: r + i for i, name in enumerate(lines)}
    labels = {
        "revenue": "Revenue (build, if it succeeds)", "gross_margin": "Gross margin %", "ap_rate": "A&P % of revenue",
        "gross_profit": "Gross profit", "marketing": "Marketing (A&P)", "rnd": "R&D", "org_cost": "Org cost",
        "other_opex_pre": "Distribution & overheads (pre-scale rule)", "ebitda_pre": "EBITDA (pre-scale rule)",
        "other_opex": "Distribution & overheads (plug after scale)", "ebitda": "EBITDA",
        "wc": "Net working capital", "wc_investment": "Working-capital investment", "fcf": "Free cash flow",
        "cash_burn": "Cash burn", "cumulative_cash": "Cumulative cash", "discount_factor": "Discount factor",
        "bought_revenue": "Bought brand revenue (Minimalist)", "bought_ebitda": "Bought brand EBITDA"}
    T, cg, ramp, m0 = sc["T"], P["category_growth"], C["ramp"], C["m0"]
    mstart_row = R["bought_ebitda"] + 1
    mstart = f"$C${mstart_row}"
    for name in lines:
        ws.cell(R[name], 2, labels[name])
        for t in YEARS:
            y = ycol[t]
            yr = f"{y}${row_year}"
            prev = ycol[t - 1] if t else None
            rv = f"{y}{R['revenue']}"
            if name == "revenue":
                post = (f"{prev}{R['revenue']}*(1+IF({yr}-{T}=1,{C['psg1']},IF({yr}-{T}=2,{C['psg2']},{cg})))"
                        if t else "0")
                f = f"=IF({yr}=0,0,IF({yr}<={T},{P['y1_revenue']}*{sc['g']}^({yr}-1),{post}))"
            elif name == "gross_margin":
                f = (f"=IF({yr}=0,0,IF({yr}<={T},IF({yr}<=2,{P['gross_margin_early']},{P['gross_margin_early']}"
                     f"+({yr}-2)/MAX({T}-2,1)*({P['gross_margin_scale']}-{P['gross_margin_early']})),"
                     f"{P['gross_margin_scale']}))")
            elif name == "ap_rate":
                f = (f"=IF({yr}=0,0,IF({yr}<={T},IF({yr}<=2,{P['ap_pct_launch']},{P['ap_pct_growth']}),"
                     f"{P['ap_pct_scale']}))")
            elif name == "gross_profit":
                f = f"={y}{R['gross_margin']}*{rv}"
            elif name == "marketing":
                f = f"={y}{R['ap_rate']}*{rv}"
            elif name == "rnd":
                f = f"={P['rnd_floor']}+{P['rnd_pct']}*{rv}"
            elif name == "org_cost":
                f = f"=IF({yr}=0,0.5*{P['org_cost_floor']},MAX({P['org_cost_floor']},{P['org_cost_pct']}*{rv}))"
            elif name == "other_opex_pre":
                f = f"=IF({yr}=0,0,{P['var_opex_pct']}*{rv})"
            elif name == "ebitda_pre":
                f = (f"={y}{R['gross_profit']}-{y}{R['marketing']}-{y}{R['other_opex_pre']}-{y}{R['org_cost']}"
                     f"-{y}{R['rnd']}")
            elif name == "ebitda":
                f = (f"=IF({yr}<={T},{y}{R['ebitda_pre']},({mstart}+({P['steady_ebitda_margin']}-{mstart})"
                     f"*MIN(MAX({yr}-{T},0)/{ramp},1))*{rv})")
            elif name == "other_opex":
                f = (f"=IF({yr}<={T},{y}{R['other_opex_pre']},{y}{R['gross_profit']}-{y}{R['marketing']}"
                     f"-{y}{R['org_cost']}-{y}{R['rnd']}-{y}{R['ebitda']})")
            elif name == "wc":
                f = f"={P['wc_days']}/365*{rv}"
            elif name == "wc_investment":
                f = f"={y}{R['wc']}-{prev}{R['wc']}" if t else f"={y}{R['wc']}"
            elif name == "fcf":
                f = f"={y}{R['ebitda']}-{y}{R['wc_investment']}"
            elif name == "cash_burn":
                f = f"=MAX(-{y}{R['fcf']},0)"
            elif name == "cumulative_cash":
                f = f"={prev}{R['cumulative_cash']}+{y}{R['fcf']}" if t else f"={y}{R['fcf']}"
            elif name == "discount_factor":
                f = f"=(1+{P['discount_rate']})^(-{yr})"
            elif name == "bought_revenue":
                f = f"={P['acquired_rev_y0']}*(1+{cg})^{yr}"
            elif name == "bought_ebitda":
                f = (f"=({m0}+({P['steady_ebitda_margin']}-{m0})*MIN({yr}/{ramp},1))*{y}{R['bought_revenue']}")
            c = ws.cell(R[name], FIRST_YC + t, f)
            c.number_format = PCT if name in ("gross_margin", "ap_rate") else ("0.000" if name == "discount_factor"
                                                                              else '#,##0.0')
    # margin at scale (start of the post-scale ramp): never below M0. Uses ebitda_pre to avoid a circular ref.
    ws.cell(mstart_row, 2, "EBITDA margin at scale year (start of ramp)")
    ws.cell(mstart_row, 3, f"=MAX({m0},INDEX({yfirst}{R['ebitda_pre']}:{ylast}{R['ebitda_pre']},{T}+1)"
                           f"/INDEX({yfirst}{R['revenue']}:{ylast}{R['revenue']},{T}+1))").number_format = PCT
    r = mstart_row + 2

    def rr(name):
        return f"${yfirst}${R[name]}:${ylast}${R[name]}"
    years = f"${yfirst}${row_year}:${ylast}${row_year}"
    df10 = f"${ylast}${R['discount_factor']}"
    disc, tg, p = P["discount_rate"], P["terminal_growth"], P["p_success"]
    S = {}
    summary = [
        ("pv_burn_to_scale", "PV of cash burn to scale (Y0..T)",
         f"=SUMPRODUCT({rr('cash_burn')},{rr('discount_factor')},({years}<={T})*1)"),
        ("peak_cumulative_burn", "Peak cumulative burn", f"=-MIN({rr('cumulative_cash')})"),
        ("p_success", "P(success)", f"={p}"),
        ("risk_adjusted_build_cost", "Risk-adjusted build cost = PV burn / p", "={pv_burn_to_scale}/{p_success}"),
        ("forgone_profits", "Forgone profits (bought EBITDA minus build EBITDA after scale, PV Y1-Y10)",
         f"=SUMPRODUCT({rr('bought_ebitda')}-{rr('ebitda')}*({years}>{T}),{rr('discount_factor')},({years}>=1)*1)"),
        ("y10_revenue_gap", "Y10 revenue gap (bought minus build)",
         f"=${ylast}${R['bought_revenue']}-${ylast}${R['revenue']}"),
        ("share_loss", "Permanent share loss (Y10 gap capitalised, PV)",
         f"={{y10_revenue_gap}}*{P['steady_ebitda_margin']}/({disc}-{tg})*{df10}"),
        ("cost_of_delay", "Cost of delay = forgone profits + share loss", "={forgone_profits}+{share_loss}"),
        ("build_equivalent_cost", "Build-equivalent cost", "={risk_adjusted_build_cost}+{cost_of_delay}"),
        ("other_buy_costs", "Other buy costs (last 9.5% + transaction + integration - cash acquired)",
         f"={P['remaining_stake_cost']}+{P['transaction_costs']}+{P['integration_cost']}-{P['cash_acquired']}"),
        ("buy_all_in_100pct", "Buy all-in (100%)", f"={C['headline_price']}+{{other_buy_costs}}"),
        ("breakeven_headline_price", "BREAKEVEN headline price for 90.5%", "={build_equivalent_cost}-{other_buy_costs}"),
        ("headroom_vs_price", "Headroom vs ₹2,706 Cr paid", f"={{breakeven_headline_price}}-{C['headline_price']}"),
        ("pv_failed_attempt", "PV of a failed attempt (Y0..k)",
         f"=SUMPRODUCT({rr('cash_burn')},{rr('discount_factor')},({years}<={sc['k']})*1)"),
        ("refined_build_cost", "Refined build cost = PV burn + (1/p - 1) x PV failed attempt",
         "={pv_burn_to_scale}+(1/{p_success}-1)*{pv_failed_attempt}"),
        ("refined_breakeven", "Refined breakeven", "={refined_build_cost}+{cost_of_delay}-{other_buy_costs}"),
        ("bought_brand_dcf_100pct", "Cross-check: bought-brand DCF (100%, EBITDA proxy)",
         f"=SUMPRODUCT({rr('bought_ebitda')},{rr('discount_factor')},({years}>=1)*1)"
         f"+${ylast}${R['bought_ebitda']}*(1+{tg})/({disc}-{tg})*{df10}"),
    ]
    for key, lab, f in summary:
        S[key] = f"$C${r}"
        ws.cell(r, 2, lab)
        c = ws.cell(r, 3, f.format(**{k: v for k, v in S.items()}) if "{" in f else f)
        c.number_format = "0.0%" if key == "p_success" else MONEY0
        c.fill = OUT_FILL
        if key == "breakeven_headline_price":
            ws.cell(r, 2).font = BOLD
            c.font = BOLD
        r += 1
    S["T"] = sc["T"]
    return S, r + 1


def build_model_tab(ws, refs):
    ws.cell(1, 1, "Model: what would it cost HUL to build Minimalist? Live model (all inputs on Assumptions)").font = \
        Font(bold=True, size=13)
    ws.cell(2, 1, "₹ Cr, pre-tax, Y0 = FY26. Build revenue is conditional on success; P(success) enters via "
                  "risk-adjusted cost. Mirrors scripts/build_model.py.")
    out = {}
    row = 4
    for name in ("Base", "Bull", "Bear"):
        S, row = model_block(ws, row, f"Scenario: {name}", refs[name], refs["const"])
        out[name] = S
    widths(ws, [2, 62] + [11] * 11)
    ws.freeze_panes = "C4"
    return out


# ---------------------------------------------------------------- sensitivity
def sensitivity_tab(ws, refs, model, mo):
    base = refs["Base"]
    C = refs["const"]
    ws.cell(1, 1, "Sensitivity: breakeven headline price for 90.5% (₹ Cr)").font = Font(bold=True, size=13)
    # key outputs summary (live links into Build_Model)
    ws.cell(3, 1, "Key outputs (live, from Build_Model)").font = BOLD
    ws.cell(3, 1).fill = SEC_FILL
    keys = [("years_from_decision", None), ("pv_burn_to_scale", MONEY0), ("p_success", PCT),
            ("risk_adjusted_build_cost", MONEY0), ("cost_of_delay", MONEY0), ("build_equivalent_cost", MONEY0),
            ("buy_all_in_100pct", MONEY0), ("breakeven_headline_price", MONEY0), ("headroom_vs_price", MONEY0)]
    for j, s in enumerate(("Base", "Bull", "Bear"), 2):
        ws.cell(4, j, s).font = BOLD
    for i, (k, fmt) in enumerate(keys, 5):
        ws.cell(i, 1, k)
        for j, s in enumerate(("Base", "Bull", "Bear"), 2):
            f = f"=Build_Model!{model[s]['T']}+1" if k == "years_from_decision" else f"=Build_Model!{model[s][k]}"
            c = ws.cell(i, j, f)
            c.number_format = fmt or "0"
            c.fill = OUT_FILL
    ws.cell(14, 1, "Price paid at completion (90.5%)")
    ws.cell(14, 2, f"={C['headline_price']}").number_format = MONEY0
    # grid
    g0 = 17
    ws.cell(g0 - 1, 1, "Grid: P(success) (rows) x years from launch to ₹347 Cr (columns); other drivers at base. "
                       "Red = breakeven above the price paid, i.e. buying was cheaper.").font = BOLD
    ws.cell(g0, 1, "P(success) \\ years")
    ps = [0.05, 0.08, 0.11, 0.15, 0.20, 0.29]
    Ts = [3, 4, 5, 6, 7, 8, 9]
    # one live base-case block per T, below the grid
    blk_row = g0 + len(ps) + 4
    ws.cell(blk_row - 1, 1, "Calculation blocks for the grid (base drivers, years-to-scale overridden by the "
                            "column header above)").font = BOLD
    t_out = {}
    for j, T in enumerate(Ts):
        hc = ws.cell(g0, 2 + j, T)
        hc.font, hc.number_format = Font(bold=True, color="0000FF"), '0" yrs"'
        P = dict(base, years_to_scale=f"Sensitivity!${L(2 + j)}${g0}")
        S, blk_row = model_block(ws, blk_row, f"Grid block: years to scale = column {L(2 + j)} header", P, C)
        t_out[T] = S
    for i, p in enumerate(ps, g0 + 1):
        pc = ws.cell(i, 1, p)
        pc.font, pc.number_format = BLUE, "0%"
        for j, T in enumerate(Ts):
            S = t_out[T]
            c = ws.cell(i, 2 + j, f"={S['pv_burn_to_scale']}/$A{i}+{S['cost_of_delay']}-{S['other_buy_costs']}")
            c.number_format = '#,##0'
    rng = f"B{g0 + 1}:{L(1 + len(Ts))}{g0 + len(ps)}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"B{g0 + 1}>$B$14"], fill=RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"B{g0 + 1}<=$B$14"], fill=GREEN))
    # static outputs from Python: tornado + Monte Carlo, placed to the right
    c0 = 16  # column P: clear of the grid calculation blocks (A..M)
    ws.cell(3, c0, "Static outputs from scripts/build_model.py (not live)").font = BOLD
    ws.cell(3, c0).fill = SEC_FILL
    ws.cell(4, c0, "Tornado: breakeven with one driver at low / high").font = BOLD
    for j, h in enumerate(("driver", "at low", "at high", "swing"), c0):
        ws.cell(5, j, h).font = BOLD
    torn = {}
    for r in mo:
        if r["section"] == "tornado":
            k, kind = re.match(r"(.+)_(at_low|at_high|swing)$", r["metric"]).groups()
            torn.setdefault(k, {})[kind] = float(r["value"])
    for i, (k, d) in enumerate(sorted(torn.items(), key=lambda x: -x[1]["swing"]), 6):
        ws.cell(i, c0, k)
        for j, kind in enumerate(("at_low", "at_high", "swing"), c0 + 1):
            ws.cell(i, j, round(d[kind], 1)).number_format = MONEY0
    mrow = 6 + len(torn) + 1
    ws.cell(mrow, c0, "Monte Carlo (10,000 triangular draws, seed 42)").font = BOLD
    for i, r in enumerate([r for r in mo if r["section"] == "monte_carlo"], mrow + 1):
        ws.cell(i, c0, r["metric"])
        v = float(r["value"])
        ws.cell(i, c0 + 1, v).number_format = PCT if r["unit"] == "ratio" else ("#,##0" if r["unit"] == "count"
                                                                                 else MONEY0)
    widths(ws, [34] + [12] * 14 + [26, 12, 12, 12])
    ws.freeze_panes = "B5"
    return {"grid": (g0, ps, Ts)}


# ---------------------------------------------------------------- conflicts (markdown -> rows)
def conflicts_tab(ws):
    text = (DATA / "conflicts.md").read_text(encoding="utf-8").splitlines()
    r = 1
    for line in text:
        s = line.strip()
        if not s or s == "---" or re.fullmatch(r"\|[-| :]+\|", s):
            continue
        if s.startswith("|"):
            cells = [c.strip().replace("**", "") for c in s.strip("|").split("|")]
            is_hdr = cells and cells[0] in ("Figure", "Field", "Item", "Source A", "Brand")
            for j, c in enumerate(cells, 1):
                cell = ws.cell(r, j + 1, c)
                cell.alignment = WRAP
                if is_hdr:
                    cell.font, cell.fill = HDR_FONT, HDR_FILL
        else:
            heading = s.startswith("#")
            cell = ws.cell(r, 1, s.lstrip("#").strip().replace("**", ""))
            if heading:
                cell.font = Font(bold=True, size=12 if s.startswith("# ") else 11)
                cell.fill = SEC_FILL
        r += 1
    widths(ws, [30, 30, 45, 55, 40])
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------- README
def readme_tab(ws):
    lines = [
        ("Buy or Build: consumer brands study (Priyanshi Paul)", "title"),
        ("Indian listed FMCG / consumer parents, 2020 onward: acquired digital-first brands (Tab A) vs in-house "
         "builds (Tab B). All money in ₹ Cr; dates ISO; FY24 = Apr 2023 to Mar 2024.", None),
        ("", None),
        ("Tabs", "h"),
        ("Tab_A_Buy: one row per acquisition; revenue_multiple is a live formula = consideration / (stake/100) / "
         "revenue (100% basis).", None),
        ("Tab_B_Build: one row per in-house brand build (core = launched 2020+, calibration = older builds used "
         "for base rates).", None),
        ("Vocabulary: controlled values and inclusion rules for category, what_was_acquired, deal_structure, "
         "row_type, base_rate_flag.", None),
        ("Assumptions: every model driver (low / base / high) with source IDs; Bull/Base/Bear columns "
         "are formulas. Edit blue cells only.", None),
        ("Build_Model: live 10-year build P&L, cost of delay and breakeven price for Base, Bull and Bear.", None),
        ("Sensitivity: live key-output summary and P(success) x years-to-scale grid; static tornado and Monte "
         "Carlo results from Python.", None),
        ("Source_Log: every source used (ID, URL, publisher, tier, dates). Raw text saved as "
         "sources/raw/S###_*.txt.", None),
        ("Conflicts: where sources disagree, each figure, its source and what it measures, plus the verdict.",
         None),
        ("Claims_Log: every AI-recalled figure, the verified value, the source and how a mismatch was caught.",
         None),
        ("AR_Mentions: brand-mention counts in parents' annual reports, FY20-FY26 (status tracking for builds).",
         None),
        ("Base_Rates: build success rates from Tab B (Wilson 95% intervals) that feed P(success).", None),
        ("Framework_Scores: Framework back-test, each row's V1-V5 scores, the predicted side and whether it was "
         "right.", None),
        ("", None),
        ("Source tiers", "h"),
        ("T1: BSE/NSE filings, annual reports, CCI orders, MCA filings, company press releases and investor "
         "presentations.", None),
        ("T2: earnings-call transcripts, DRHP/RHP, reputable press directly quoting company officials.", None),
        ("T3: Inc42, Entrackr, ET, Mint, Moneycontrol, Business Standard, Tracxn/Crunchbase summaries.", None),
        ("T4: blogs, Wikipedia, AI memory. Never enough on its own.", None),
        ("", None),
        ("Confidence rules", "h"),
        ("High (green): a T1 source with a verbatim quote.", "High"),
        ("Medium (amber): a T2 source, or two independent T3 sources that agree.", "Medium"),
        ("Low (red): a single T3 source, or a derived or inferred value. UNCONFIRMED (red): value left blank, "
         "with the reason in notes; never estimated silently.", "Low"),
        ("", None),
        ("Conventions", "h"),
        ("Blue font = hard-coded input. Yellow fill = formula output. Revenue basis is recorded per row "
         "(revenue_from_operations vs total_income). Conflicts resolved in favour of the most authoritative "
         "source.", None),
    ]
    fills = {"High": GREEN, "Medium": AMBER, "Low": RED}
    for i, (t, kind) in enumerate(lines, 1):
        c = ws.cell(i, 1, t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            c.font = Font(bold=True, size=14)
        elif kind == "h":
            c.font, c.fill = BOLD, SEC_FILL
        elif kind in fills:
            c.fill = fills[kind]
    ws.column_dimensions["A"].width = 130


# ---------------------------------------------------------------- main
def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    readme_tab(ws)

    deal_tab(wb.create_sheet("Tab_A_Buy"), read("tab_a.csv"))
    deal_tab(wb.create_sheet("Tab_B_Build"), read("tab_b.csv"))
    dump(wb.create_sheet("Vocabulary"), read("vocabulary.csv"), wide={"inclusion_rule"})
    refs = assumptions_tab(wb.create_sheet("Assumptions"), read("assumptions.csv"))
    model = build_model_tab(wb.create_sheet("Build_Model"), refs)
    sensitivity_tab(wb.create_sheet("Sensitivity"), refs, model, read("model_outputs.csv"))
    dump(wb.create_sheet("Source_Log"), read("source_log.csv"), wide={"title", "notes", "url"}, url_cols={"url"})
    conflicts_tab(wb.create_sheet("Conflicts"))
    ai = wb.create_sheet("Claims_Log")
    cols = dump(ai, read("claims_verification_log.csv"), wide={"item", "value_recalled", "value_verified", "how_caught"})
    mcol = L(cols.index("match") + 1)
    n = ai.max_row
    ai.conditional_formatting.add(f"{mcol}2:{mcol}{n}", FormulaRule(formula=[f'{mcol}2="Y"'], fill=GREEN))
    ai.conditional_formatting.add(f"{mcol}2:{mcol}{n}", FormulaRule(formula=[f'{mcol}2="Partial"'], fill=AMBER))
    ai.conditional_formatting.add(f"{mcol}2:{mcol}{n}", FormulaRule(formula=[f'{mcol}2="N"'], fill=RED))
    dump(wb.create_sheet("AR_Mentions"), read("ar_mentions.csv"), wide={"sample_quote"}, url_cols={"ar_url"})
    dump(wb.create_sheet("Base_Rates"), read("base_rates.csv"), wide={"definition", "rows_included", "notes"})
    dump(wb.create_sheet("Framework_Scores"), read("framework_scores.csv"), wide={"brand"})

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Saved {OUT} ({len(wb.sheetnames)} tabs)")


if __name__ == "__main__":
    main()
